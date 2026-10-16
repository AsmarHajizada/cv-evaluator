from io import BytesIO

import docx
import pytest

from cv_eval.extraction import ExtractionError, extract_cv
from cv_eval.schemas import PassageKind


def _by_text(cv, fragment):
    return next(p for p in cv.passages if fragment in p.text)


def test_markdown_keeps_section_hierarchy_and_inline_formatting(sample_cv):
    bullet = _by_text(sample_cv, "demand forecasting model")
    assert bullet.kind == PassageKind.bullet
    assert bullet.section.startswith("Experience › Machine Learning Engineer")

    # "**Contract Q&A assistant** (personal project): ..." stays one passage
    project = _by_text(sample_cv, "Contract Q&A assistant")
    assert "retrieval-augmented generation" in project.text
    assert project.section == "Projects"

    assert _by_text(sample_cv, "Kubernetes").section == "Skills"
    assert [p.id for p in sample_cv.passages] == [
        f"P{i}" for i in range(1, len(sample_cv.passages) + 1)
    ]


def test_plain_text_lines_become_separate_passages():
    text = (
        "Jane Doe\njane@example.com\n\nEXPERIENCE\n• Built X\n• Shipped Y\nSkills:\nPython, SQL\n"
    )
    cv = extract_cv(text.encode(), "cv.txt")
    texts = [(p.section, p.kind, p.text) for p in cv.passages]
    assert texts == [
        (None, PassageKind.paragraph, "Jane Doe"),
        (None, PassageKind.paragraph, "jane@example.com"),
        ("Experience", PassageKind.bullet, "Built X"),
        ("Experience", PassageKind.bullet, "Shipped Y"),
        ("Skills", PassageKind.paragraph, "Python, SQL"),
    ]


def test_docx_headings_bullets_and_tables():
    d = docx.Document()
    d.add_heading("Experience", level=1)
    d.add_paragraph("Deployed a fraud model to production.", style="List Bullet")
    d.add_heading("Skills", level=1)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Languages", "Python, SQL"
    table.cell(1, 0).text, table.cell(1, 1).text = "Cloud", "AWS"
    buf = BytesIO()
    d.save(buf)

    cv = extract_cv(buf.getvalue(), "cv.docx")
    bullet = _by_text(cv, "fraud model")
    assert (bullet.section, bullet.kind) == ("Experience", PassageKind.bullet)
    rows = [p.text for p in cv.passages if p.kind == PassageKind.table_row]
    assert rows == ["Languages | Python, SQL", "Cloud | AWS"]
    assert all(p.section == "Skills" for p in cv.passages if p.kind == PassageKind.table_row)


def test_unsupported_and_empty_files_raise():
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_cv(b"x", "cv.pages")
    with pytest.raises(ExtractionError, match="No text"):
        extract_cv(b"   \n", "cv.txt")
